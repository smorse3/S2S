from openpyxl import Workbook, load_workbook


class SpreadsheetHandler:

    def __init__(self):

        self.workbook = None
        self.sheet = None

        self.current_path = None

        self.current_row = 1
        self.current_col = 1

    ###################################################
    # FILES
    ###################################################

    def create_new(self, path):

        self.workbook = Workbook()

        self.sheet = self.workbook.active

        self.current_path = path

        self.save()

    def load(self, path):

        self.workbook = load_workbook(path)

        self.sheet = self.workbook.active

        self.current_path = path

    def save(self):

        if self.current_path:

            self.workbook.save(self.current_path)

    ###################################################
    # CELL WRITING
    ###################################################

    def write_current_cell(self, text):

        self.sheet.cell(
            row=self.current_row,
            column=self.current_col,
            value=text
        )

        self.save()

    ###################################################
    # NAVIGATION
    ###################################################

    def move_right(self):

        self.current_col += 1

    def move_left(self):

        if self.current_col > 1:
            self.current_col -= 1

    def move_down(self):

        self.current_row += 1

    def move_up(self):

        if self.current_row > 1:
            self.current_row -= 1

    def get_position(self):

        return (
            self.current_row,
            self.current_col
        )