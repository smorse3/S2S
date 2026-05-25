import pandas as pd

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import CellIsRule

from spreadsheet.conditional_formatting import ConditionalFormattingRule


class SpreadsheetIO:

    # =========================================================
    # SAVE XLSX (DATA + CONDITIONAL FORMATTING + RULE METADATA)
    # =========================================================
    @staticmethod
    def save_xlsx(model, path):

        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        from openpyxl.formatting.rule import CellIsRule

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # =========================================
        # DETERMINE GRID SIZE SAFELY
        # =========================================
        try:
            df = model.df
            max_rows = max(len(df.index), 1)
            max_cols = max(len(df.columns), 1)
        except Exception:
            max_rows = 1
            max_cols = 1
            df = model.df

        # =========================================
        # WRITE ALL CELL VALUES (SAFE GRID SCAN)
        # =========================================
        for r in range(max_rows):
            for c in range(max_cols):

                try:
                    value = model.get_cell(r, c)
                except Exception:
                    value = None

                ws.cell(
                    row=r + 1,
                    column=c + 1,
                    value=value
                )

        # =========================================
        # APPLY EXCEL CONDITIONAL FORMATTING
        # =========================================
        for rule in model.conditional_rules:

            try:

                fill = PatternFill(
                    start_color=rule.color.replace("#", ""),
                    end_color=rule.color.replace("#", ""),
                    fill_type="solid"
                )

                if rule.rule_type == "between":

                    ws.conditional_formatting.add(
                        rule.target_range,
                        CellIsRule(
                            operator="between",
                            formula=[
                                str(rule.min_value),
                                str(rule.max_value)
                            ],
                            fill=fill
                        )
                    )

                elif rule.rule_type == "not_between":

                    ws.conditional_formatting.add(
                        rule.target_range,
                        CellIsRule(
                            operator="notBetween",
                            formula=[
                                str(rule.min_value),
                                str(rule.max_value)
                            ],
                            fill=fill
                        )
                    )

                elif rule.rule_type == "greater_than":

                    ws.conditional_formatting.add(
                        rule.target_range,
                        CellIsRule(
                            operator="greaterThan",
                            formula=[str(rule.min_value)],
                            fill=fill
                        )
                    )

                elif rule.rule_type == "less_than":

                    ws.conditional_formatting.add(
                        rule.target_range,
                        CellIsRule(
                            operator="lessThan",
                            formula=[str(rule.min_value)],
                            fill=fill
                        )
                    )

                elif rule.rule_type == "equal":

                    ws.conditional_formatting.add(
                        rule.target_range,
                        CellIsRule(
                            operator="equal",
                            formula=[str(rule.value)],
                            fill=fill
                        )
                    )

            except Exception as e:
                print("CF export error:", e)

        # =========================================
        # WRITE RULE METADATA (HIDDEN SHEET)
        # =========================================
        rules_ws = wb.create_sheet("__S2S_RULES__")
        rules_ws.sheet_state = "hidden"

        headers = [
            "rule_type",
            "min_value",
            "max_value",
            "value",
            "color",
            "target_range",
            "stop_if_true"
        ]

        for col, h in enumerate(headers, start=1):
            rules_ws.cell(row=1, column=col, value=h)

        for i, rule in enumerate(model.conditional_rules, start=2):

            rules_ws.cell(i, 1, rule.rule_type)
            rules_ws.cell(i, 2, rule.min_value)
            rules_ws.cell(i, 3, rule.max_value)
            rules_ws.cell(i, 4, rule.value)
            rules_ws.cell(i, 5, rule.color)
            rules_ws.cell(i, 6, rule.target_range)
            rules_ws.cell(i, 7, rule.stop_if_true)

        wb.save(path)

    # =========================================================
    # LOAD XLSX (DATA + RULES)
    # =========================================================
    @staticmethod
    def load(path):

        wb = load_workbook(path)

        # ================================
        # LOAD SHEET CLEANLY (NO PAD SHIFT)
        # ================================
        ws = wb["Sheet1"]

        data = []

        for row in ws.iter_rows(values_only=True):

            # skip fully empty rows
            if all(cell is None for cell in row):
                continue

            data.append(list(row))

        # normalize rectangular grid
        max_cols = max(len(r) for r in data) if data else 1

        normalized = []
        for r in data:
            r = list(r)
            r += [None] * (max_cols - len(r))
            normalized.append(r)

        df = pd.DataFrame(normalized)

        # -----------------------------
        # LOAD CONDITIONAL RULES
        # -----------------------------
        rules = []

        if "__S2S_RULES__" in wb.sheetnames:

            rules_ws = wb["__S2S_RULES__"]

            for row in rules_ws.iter_rows(min_row=2, values_only=True):

                try:
                    if not row:
                        continue

                    rule = ConditionalFormattingRule(

                        rule_type=row[0] if len(row) > 0 else None,
                        min_value=row[1] if len(row) > 1 else None,
                        max_value=row[2] if len(row) > 2 else None,
                        value=row[3] if len(row) > 3 else None,
                        color=row[4] if len(row) > 4 else "#FFFF00",
                        target_range=row[5] if len(row) > 5 else "A1:Z100",
                        stop_if_true=bool(row[6]) if len(row) > 6 else False
                    )

                    rules.append(rule)

                except Exception as e:
                    print("Rule load failed:", e)

        return df, rules